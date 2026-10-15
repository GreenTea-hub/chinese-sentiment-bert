# -*- coding: utf-8 -*-
"""Convert the opening-report Markdown into a formatted .docx."""
import re
import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


SRC = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "docs",
                   "开题报告-自然语言处理.md")
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "docs",
                   "开题报告-自然语言处理.docx")


def set_run_font(run, east="宋体", latin="Times New Roman", size=12, bold=None,
                 color=None):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), east)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_runs_with_bold(paragraph, text, east="宋体", latin="Times New Roman",
                       size=12, color=None):
    """Split text on **bold** markers and add runs."""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, east=east, latin=latin, size=size, bold=True,
                         color=color)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, east=east, latin=latin, size=size, bold=False,
                         color=color)


def shade_cell(cell, color="D9E2F3"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    add_runs_with_bold(p, text, size=size)
    for run in p.runs:
        run.font.bold = bold
        run.font.size = Pt(size)


def add_table(doc, rows):
    # rows: list of list of str
    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(n_cols):
            text = row[j] if j < len(row) else ""
            set_cell_text(table.cell(i, j), text, bold=(i == 0))
            if i == 0:
                shade_cell(table.cell(i, j))
    doc.add_paragraph()


def main():
    with open(SRC, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()

    # page margins
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(3.0)

    # default style
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    heading_font_map = {
        1: ("黑体", 16),
        2: ("黑体", 14),
        3: ("黑体", 12),
    }

    i = 0
    in_table = False
    table_rows = []
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # table
        if stripped.startswith("|"):
            if not in_table:
                in_table = True
                table_rows = []
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            # skip separator rows like | --- | --- |
            if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                table_rows.append(cells)
            i += 1
            continue
        if in_table:
            in_table = False
            add_table(doc, table_rows)
            table_rows = []

        # headings
        m = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            east, size = heading_font_map.get(level, ("黑体", 12))
            p = doc.add_heading("", level=level)
            add_runs_with_bold(p, text, east=east, latin="Arial", size=size)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
                run.font.bold = True
                run.font.name = "Arial"
                rpr = run._element.get_or_add_rPr()
                rfonts = rpr.find(qn("w:rFonts"))
                if rfonts is None:
                    rfonts = OxmlElement("w:rFonts")
                    rpr.append(rfonts)
                rfonts.set(qn("w:eastAsia"), east)
            i += 1
            continue

        # blockquote
        if stripped.startswith(">"):
            text = stripped.lstrip(">").strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.74)
            add_runs_with_bold(p, text, size=10.5)
            for run in p.runs:
                run.font.italic = True
            i += 1
            continue

        # bullet / numbered list
        m = re.match(r"^[-*]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_runs_with_bold(p, m.group(1))
            i += 1
            continue
        m = re.match(r"^\d+[.、]\s+(.*)$", stripped)
        if m:
            # render the number literally so separate lists do not continue
            # Word's shared numbering sequence
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.74)
            p.paragraph_format.first_line_indent = Cm(-0.74)
            p.paragraph_format.line_spacing = 1.5
            add_runs_with_bold(p, stripped)
            i += 1
            continue

        # normal paragraph
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)
        p.paragraph_format.line_spacing = 1.5
        add_runs_with_bold(p, stripped)
        i += 1

    if in_table and table_rows:
        add_table(doc, table_rows)

    doc.save(OUT)
    print("saved:", OUT)


if __name__ == "__main__":
    sys.exit(main())
